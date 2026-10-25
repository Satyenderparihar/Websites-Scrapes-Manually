import openpyxl
import os
import sys
import re
import urllib.parse

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\All In One Programming Resources.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)
wb_data = openpyxl.load_workbook(fpath, data_only=True)

VALID_TLDS = set([
    'com', 'org', 'net', 'edu', 'gov', 'mil', 'io', 'ai', 'co', 'app', 'dev', 'tv', 'me', 'in', 'is', 'so', 'gg',
    'cc', 'ly', 'to', 'la', 'fm', 'xyz', 'tech', 'online', 'site', 'store', 'club', 'pro', 'studio', 'tools', 'garden',
    'lol', 'live', 'page', 'space', 'fun', 'design', 'world', 'media', 'stream', 'cloud', 'engineer', 'academy',
    'software', 'download', 'community', 'systems', 'education', 'solutions', 'technology', 'training'
])

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip().strip("'\"<>[]()")
    if " " in u or "\n" in u or "\t" in u:
        return ""
    if not u.startswith("http://") and not u.startswith("https://"):
        if "." in u and not u.startswith("="):
            u = "https://" + u
    parsed = urllib.parse.urlparse(u)
    if not parsed.netloc:
        return ""
    parts = parsed.netloc.lower().split('.')
    if len(parts) < 2:
        return ""
    tld = parts[-1].split(':')[0]
    if tld not in VALID_TLDS and len(tld) > 10:
        return ""
    return u.rstrip('/')

def get_domain_key(url):
    parsed = urllib.parse.urlparse(url)
    domain = parsed.netloc.lower()
    if domain.startswith("www."):
        domain = domain[4:]
    path = parsed.path.rstrip('/')
    query = parsed.query
    if any(sub in domain for sub in ['youtube.com', 'instagram.com', 't.me', 'github.com', 'geeksforgeeks.org', 'reddit.com', 'edx.org', 'coursera.org', 'w3schools.com', 'coddy.tech', 'classcentral.com', 'botdownloader.com', 'analyticsvidhya.com', 'stpcomputereducation.com', 'cloudskillsboost.google', 'data-flair.training', 'docs.google.com']):
        key_suffix = path
        if query:
            key_suffix = f"{path}?{query}" if path else f"?{query}"
        if key_suffix:
            return f"{domain}{key_suffix}"
    return domain

extracted_tools = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_d = wb_data[sname]
    
    # 1. Cheatsheet-list
    if sname == "Cheatsheet-list":
        for r in range(1, ws.max_row + 1):
            name = str(ws.cell(r, 1).value or '').strip()
            cell_2 = ws.cell(r, 2)
            url = cell_2.hyperlink.target if cell_2.hyperlink else str(cell_2.value or '').strip()
            if not name or name in ["None"] or len(name) < 3:
                continue
            cu = clean_url(url)
            if cu:
                key = get_domain_key(cu)
                extracted_tools[key] = {
                    "name": name,
                    "url": cu,
                    "website_type": "Developer Reference & Cheat Sheets",
                    "category": "💻 Coding & Developer Platforms",
                    "pricing": "Free",
                    "description": f"Comprehensive developer cheat sheet and reference spreadsheet: {name}",
                    "source": f"Excel: All In One Programming Resources.xlsx ({sname})"
                }
        continue

    # 2. Stay updated on Telegram
    if sname == "Stay updated on Telegram":
        for r in range(1, ws.max_row + 1):
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2_cell = ws.cell(r, 2)
            c3 = str(ws.cell(r, 3).value or '').strip()
            if not c1 or any(k in c1 for k in ["NOTE:", "Channel / Group Name", "None"]) or len(c1) < 3:
                continue
            url = c2_cell.hyperlink.target if c2_cell.hyperlink else str(c2_cell.value or '').strip()
            cu = clean_url(url)
            if cu:
                key = get_domain_key(cu)
                desc = c3 if c3 and len(c3) > 10 else f"Technical developer community on Telegram: {c1}"
                extracted_tools[key] = {
                    "name": f"{c1} (Telegram)",
                    "url": cu,
                    "website_type": "Telegram Tech Community",
                    "category": "💻 Coding & Developer Platforms",
                    "pricing": "Free",
                    "description": desc,
                    "source": f"Excel: All In One Programming Resources.xlsx ({sname})"
                }
        continue

    # 3. Learn Python, Learn Machine Learning, Learn Data Science, Learn R
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell_d = ws_d.cell(r, c)
            val = str(cell_d.value or cell.value or '').strip()
            hl = cell.hyperlink.target if cell.hyperlink else None
            url = hl if hl else (val if val.startswith("http") else "")
            
            if not val or val in ["None", "Link", "Click Here"] or len(val) < 2:
                continue
            if val.startswith("NOTE:") or val.startswith("http"):
                # if val is url, name might be derived or in hl
                val_name = ""
            else:
                val_name = val
                
            cu = clean_url(url)
            if cu and "chatgpt.com" not in cu:
                key = get_domain_key(cu)
                
                # Determine name
                name = val_name if val_name else cu.split('/')[-1].replace('-', ' ').title()
                
                # Determine website type based on sheet and name
                wtype = "Developer Reference & Cheat Sheets"
                p_lower = name.lower()
                
                if "project" in p_lower:
                    if "python" in sname.lower() or "python" in p_lower:
                        wtype = "Python Project Implementation"
                    elif "machine learning" in sname.lower() or "ml" in p_lower:
                        wtype = "Machine Learning Project Implementation"
                    elif "data science" in sname.lower():
                        wtype = "Data Science Project Implementation"
                    elif "r" in sname.lower():
                        wtype = "R Project Implementation"
                    else:
                        wtype = "Coding Project Implementation"
                elif "learn python" in sname.lower() or "python" in p_lower:
                    wtype = "Python Programming Resource"
                elif "machine learning" in sname.lower() or any(m in p_lower for m in ['neural', 'cnn', 'rnn', 'ann', 'svm', 'clustering', 'deep learning']):
                    wtype = "Machine Learning & AI Resource"
                elif "data science" in sname.lower() or any(d in p_lower for d in ['pandas', 'numpy', 'matplotlib', 'seaborn', 'statistics', 'tableau']):
                    wtype = "Data Science & Analytics Resource"
                elif "learn r" in sname.lower() or " r " in f" {p_lower} ":
                    wtype = "R Programming Resource"
                    
                extracted_tools[key] = {
                    "name": name,
                    "url": cu,
                    "website_type": wtype,
                    "category": "💻 Coding & Developer Platforms",
                    "pricing": "Free",
                    "description": f"Verified tutorial & implementation guide: {name} ({wtype})",
                    "source": f"Excel: All In One Programming Resources.xlsx ({sname})"
                }

print(f"\nTotal Unique Resources Extracted from All In One: {len(extracted_tools)}")

# Type breakdown
t_counts = {}
for item in extracted_tools.values():
    t = item["website_type"]
    t_counts[t] = t_counts.get(t, 0) + 1

print("\nBreakdown by Website Type:")
for t, cnt in sorted(t_counts.items(), key=lambda x: x[1], reverse=True):
    print(f"  {t:45s}: {cnt}")
