import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\All In One Programming Resources.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)

for sname in ['Learn Python', 'Learn Machine Learning']:
    ws = wb[sname]
    print("=" * 60)
    print(f"SHEET: {sname}")
    print("=" * 60)
    for r in range(1, 15):
        row_items = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            val = str(cell.value or '').strip()
            hl = cell.hyperlink.target if cell.hyperlink else None
            if val or hl:
                row_items.append((c, val, hl))
        if row_items:
            print(f"Row {r:2d}: {row_items[:3]}")
