import openpyxl
import os
import sys
import re
import urllib.parse

sys.stdout.reconfigure(encoding='utf-8')

fpath = r"C:\Users\satyanarayan verma\Downloads\💻 Coding Websites.xlsx"
wb = openpyxl.load_workbook(fpath, data_only=True)
wb_f = openpyxl.load_workbook(fpath, data_only=False)

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip().strip("'\"<>[]()")
    if " " in u or "\n" in u or "\t" in u:
        return ""
    if not u.startswith("http://") and not u.startswith("https://"):
        if "." in u and not u.startswith("="):
            u = "https://" + u
    parsed = urllib.parse.urlparse(u)
    if not parsed.netloc:
        return ""
    return u.rstrip('/')

coding_records = []

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    sheet_count = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell_f = ws_f.cell(r, c)
            hl = cell.hyperlink.target if cell.hyperlink else None
            val = str(cell.value or '').strip()
            
            url = ""
            if hl:
                url = hl
            elif val.startswith("http"):
                url = val
            elif cell_f.value and str(cell_f.value).startswith("=HYPERLINK"):
                m = re.search(r'=HYPERLINK\("?([^",\)]+)"?', str(cell_f.value))
                if m:
                    url = m.group(1).strip()
            
            if url:
                # Get site name
                name = val if (val and not val.startswith("http") and val != "Link" and val != "Click Here") else ""
                if not name and c > 1:
                    prev_val = str(ws.cell(r, c-1).value or '').strip()
                    if prev_val and prev_val != "None":
                        name = prev_val
                if not name and r > 1:
                    row_vals = [str(ws.cell(r, x).value or '').strip() for x in range(1, ws.max_column+1) if ws.cell(r, x).value]
                    for rv in row_vals:
                        if rv and not rv.startswith("http") and rv != "Link" and rv != "Click Here" and len(rv) > 2:
                            name = rv
                            break
                            
                cleaned = clean_url(url)
                if cleaned:
                    coding_records.append({
                        "sheet": sname,
                        "row": r,
                        "col": c,
                        "name": name or cleaned.split('//')[-1].split('/')[0],
                        "url": cleaned
                    })
                    sheet_count += 1
    print(f"Sheet [{sname}]: Extracted {sheet_count} raw link items")

print(f"\nTotal extracted raw link items across all sheets: {len(coding_records)}")
