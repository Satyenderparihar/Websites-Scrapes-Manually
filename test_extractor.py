import os
import sys
import glob
import re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

downloads_dir = r"C:\Users\satyanarayan verma\Downloads"
files = glob.glob(os.path.join(downloads_dir, "*.xlsx"))

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip()
    if not u.startswith("http://") and not u.startswith("https://"):
        if "." in u and not u.startswith("="):
            u = "https://" + u
    return u

all_records = []

for f in files:
    fname = os.path.basename(f)
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        wb_f = openpyxl.load_workbook(f, data_only=False)
        for sname in wb.sheetnames:
            ws = wb[sname]
            ws_f = wb_f[sname]
            
            # Find header row
            header_row = 1
            headers = []
            for r in range(1, min(ws.max_row + 1, 15)):
                row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
                joined = " ".join(row_vals).lower()
                if any(k in joined for k in ['website', 'url', 'link', 'course name', 'resource name']):
                    header_row = r
                    headers = [str(ws.cell(r, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
                    break
            
            for r in range(header_row + 1, ws.max_row + 1):
                name, url, cat, pricing, desc = "", "", "", "", ""
                
                # Check hyperlinks on cells
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(r, c)
                    cell_f = ws_f.cell(r, c)
                    if cell.hyperlink and cell.hyperlink.target:
                        url = cell.hyperlink.target
                    elif cell_f.value and str(cell_f.value).startswith("=HYPERLINK"):
                        m = re.search(r'=HYPERLINK\("?([^",\)]+)"?', str(cell_f.value))
                        if m:
                            ref = m.group(1).strip()
                            if ref.startswith("http"):
                                url = ref
                            else:
                                try:
                                    url = ws[ref].value
                                except Exception:
                                    pass
                
                # Check columns by header
                for c, h in enumerate(headers, 1):
                    val = str(ws.cell(r, c).value or "").strip()
                    if not val or val == "None":
                        continue
                    if any(k in h for k in ['url', 'link']) or ('website' in h and ('.' in val or 'http' in val)):
                        if not url or ('http' not in url and 'http' in val):
                            url = val
                    elif any(k in h for k in ['name', 'resource', 'course']):
                        if not name:
                            name = val
                    elif any(k in h for k in ['type', 'category']):
                        if not cat:
                            cat = val
                    elif any(k in h for k in ['free', 'pricing', 'paid']):
                        if not pricing:
                            pricing = val
                    elif 'desc' in h:
                        if not desc:
                            desc = val
                
                # Fallback for URLs
                if not url:
                    for c in range(1, ws.max_column + 1):
                        v = str(ws.cell(r, c).value or "").strip()
                        if v.startswith("http://") or v.startswith("https://"):
                            url = v
                            break

                # If name is still empty, fallback
                if not name and url:
                    for c in range(1, ws.max_column + 1):
                        v = str(ws.cell(r, c).value or "").strip()
                        if v and v != url and v != "Click Here" and len(v) > 2:
                            name = v
                            break

                cleaned_u = clean_url(url)
                if cleaned_u and cleaned_u.startswith("http"):
                    all_records.append({
                        'source_file': fname,
                        'sheet': sname,
                        'name': name or cleaned_u.split('//')[-1].split('/')[0],
                        'url': cleaned_u,
                        'category': cat or sname,
                        'pricing': pricing or "Free",
                        'description': desc
                    })
    except Exception as e:
        print(f"Error reading {fname}: {e}")

print(f"Successfully extracted: {len(all_records)} website records from Downloads!")
for i, rec in enumerate(all_records[:15]):
    print(f"[{i+1}] {rec['name']} | {rec['url']} | {rec['category']} | {rec['pricing']}")
