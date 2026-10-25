import openpyxl
import json
import sys
import urllib.parse
import re

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\💻 Coding Websites.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
wb_f = openpyxl.load_workbook(fpath, data_only=False)

with open(r'D:\Antigravity - Agent\🌐 Websites Scrapes Manually\💾 websites_database.json', encoding='utf-8') as f:
    db = json.load(f)

db_keys = set(db.keys())
db_urls = set(v['url'].lower().rstrip('/') for v in db.values())
db_names = set(v['name'].lower().strip() for v in db.values())

print(f"Total websites in database: {len(db)}")
coding_in_db = [v for v in db.values() if v.get('category') == '💻 Coding & Developer Platforms']
print(f"Total in '💻 Coding & Developer Platforms': {len(coding_in_db)}\n")

total_raw_valid_entries = 0
sheet_audit = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    valid_entries = []
    
    if sname in ["All Project Ideas", "C Project Ideas"]:
        for r in range(1, ws.max_row + 1):
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2 = str(ws.cell(r, 2).value or '').strip()
            pname = ""
            if sname == "All Project Ideas":
                if c2 and c2 not in ["Project Title", "Link", "Click Here"] and len(c2) > 3:
                    pname = c2
                elif c1 and not any(k in c1 for k in ["Category", "Projects", "Project Ideas"]) and len(c1) > 3:
                    pname = c1
            else:
                if c1 and not any(k in c1 for k in ["Project Name", "Project Ideas", "C Project", "C++ Project", "Java Project", "Python Project"]) and len(c1) > 3:
                    pname = c1
            if pname and len(pname) > 3:
                valid_entries.append((r, pname))
                
    elif sname == "Y-I-T - BH":
        for r in range(2, ws.max_row + 1):
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2 = str(ws.cell(r, 2).value or '').strip()
            if c2 and c2 not in ["Channel Name", "Link", "None"] and len(c2) > 2:
                valid_entries.append((r, c2))
                
    elif sname == "Pythons - BH":
        for r in range(2, ws.max_row + 1):
            c2 = str(ws.cell(r, 2).value or '').strip()
            if c2 and c2 not in ["Website", "None"] and len(c2) > 2:
                valid_entries.append((r, c2))
                
    elif sname == "Sheet32 - BH":
        for r in range(2, ws.max_row + 1):
            for c in range(1, 9):
                cell = ws.cell(r, c)
                cell_f = ws_f.cell(r, c)
                val = str(cell.value or '').strip()
                hl = cell.hyperlink.target if cell.hyperlink else None
                if val and val not in ["Coding Platforms", "None"] and len(val) > 1:
                    valid_entries.append((r, val, hl))
                elif hl:
                    valid_entries.append((r, hl, hl))
                    
    else: # standard sheets
        for r in range(2, ws.max_row + 1):
            wtype = str(ws.cell(r, 1).value or '').strip()
            name = str(ws.cell(r, 2).value or '').strip()
            cell_3 = ws.cell(r, 3)
            cell_3_f = ws_f.cell(r, 3)
            url = cell_3.hyperlink.target if cell_3.hyperlink else str(cell_3.value or '').strip()
            if name and name not in ["Website Name", "None"] and len(name) > 1:
                valid_entries.append((r, name, url))
                
    sheet_audit[sname] = valid_entries
    total_raw_valid_entries += len(valid_entries)
    print(f"Sheet: {sname:35s} | Non-empty Raw Entries: {len(valid_entries)}")

print(f"\nTotal Raw Entries across all 12 sheets: {total_raw_valid_entries}")
