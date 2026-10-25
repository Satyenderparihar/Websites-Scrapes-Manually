import openpyxl
import json
import sys
import re
import urllib.parse

sys.stdout.reconfigure(encoding='utf-8')

f_course = r"C:\Users\satyanarayan verma\Downloads\🎓 Course Library.xlsx"
wb = openpyxl.load_workbook(f_course, data_only=True)
wb_f = openpyxl.load_workbook(f_course, data_only=False)

with open(r"D:\Antigravity - Agent\🌐 Websites Scrapes Manually\💾 websites_database.json", 'r', encoding='utf-8') as f:
    db = json.load(f)

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip().strip("'\"<>[]()")
    if not u.startswith("http"):
        if "." in u:
            u = "https://" + u
    parsed = urllib.parse.urlparse(u)
    clean_query = []
    if parsed.query:
        for qp in parsed.query.split("&"):
            if not any(qp.startswith(t) for t in ['utm_', 'fbclid', 'gclid', 'ref_']):
                clean_query.append(qp)
    query_str = "&".join(clean_query)
    clean = urllib.parse.urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, query_str, ""))
    return clean.rstrip('/')

def get_domain_key(url):
    parsed = urllib.parse.urlparse(url)
    domain = parsed.netloc.lower()
    if domain.startswith("www."):
        domain = domain[4:]
    path = parsed.path.rstrip('/')
    if len(path) > 1 and any(sub in domain for sub in ['botdownloader.com', 'youtube.com', 'instagram.com', 'github.com', 'coursera.org', 'analyticsvidhya.com', 'stpcomputereducation.com', 'cloudskillsboost.google']):
        return f"{domain}{path}"
    return domain

print("=== AUDITING ALL SHEETS IN 🎓 Course Library.xlsx ===")

total_all_sheets = 0
total_found_in_db = 0

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    sheet_items = []
    
    if sname == 'Sheet1':
        for r in range(1, ws.max_row + 1):
            val_a = str(ws.cell(r, 1).value or '').strip()
            if val_a.replace('.', '').isdigit() or (ws.cell(r, 2).value and 'course' not in str(ws.cell(r, 2).value).lower() and not val_a.startswith('SNo') and not val_a.startswith('Top')):
                course_name = str(ws.cell(r, 2).value or '').strip()
                if not course_name or 'Course Name' in course_name:
                    continue
                url = ""
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(r, c)
                    cell_f = ws_f.cell(r, c)
                    if cell.hyperlink and cell.hyperlink.target:
                        url = cell.hyperlink.target
                        break
                    elif cell_f.value and str(cell_f.value).startswith("=HYPERLINK"):
                        m = re.search(r'=HYPERLINK\("?([^",\)]+)"?', str(cell_f.value))
                        if m:
                            ref = m.group(1).strip()
                            if ref.startswith("http"):
                                url = ref
                            else:
                                try:
                                    url = ws[ref].value
                                except Exception:
                                    pass
                            break
                    elif cell.value and str(cell.value).startswith("http"):
                        url = str(cell.value).strip()
                        break
                if url:
                    sheet_items.append((course_name, url))
    else:
        headers = [str(ws.cell(1, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1):
            name, url = "", ""
            for c, h in enumerate(headers, 1):
                v = str(ws.cell(r, c).value or "").strip()
                if "name" in h:
                    name = v
                elif "url" in h:
                    url = v
            if url and "http" in url:
                sheet_items.append((name, url))

    found = 0
    missing = []
    for name, url in sheet_items:
        c_url = clean_url(url)
        k = get_domain_key(c_url)
        if k in db:
            found += 1
        else:
            missing.append((name, url, k))
            
    total_all_sheets += len(sheet_items)
    total_found_in_db += found
    print(f"\nSheet [{sname}]:")
    print(f"  Total items: {len(sheet_items)} | Found in DB: {found} | Missing: {len(missing)}")
    if missing:
        for m in missing:
            print(f"    ❌ MISSING: {m[0]} -> {m[1]}")
    else:
        print(f"    ✅ 100% of rows from [{sname}] are in Database!")

print("\n" + "=" * 50)
print(f"SUMMARY: Total courses in file across all 8 sheets = {total_all_sheets}")
print(f"Total courses verified in Database = {total_found_in_db}")
print("=" * 50)
