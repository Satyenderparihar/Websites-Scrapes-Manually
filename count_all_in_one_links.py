import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\All In One Programming Resources.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=False)

for sname in wb.sheetnames:
    ws = wb[sname]
    link_count = 0
    unique_links = set()
    rows_with_links = 0
    
    for r in range(1, ws.max_row + 1):
        has_link = False
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            hl = cell.hyperlink.target if cell.hyperlink else None
            val = str(cell.value or '').strip()
            u = hl if hl else (val if val.startswith('http') else '')
            if u:
                link_count += 1
                unique_links.add(u)
                has_link = True
        if has_link:
            rows_with_links += 1
            
    print(f"Sheet: {sname:30s} | Total Links: {link_count:5d} | Unique Links: {len(unique_links):5d} | Rows with Links: {rows_with_links:4d}")
