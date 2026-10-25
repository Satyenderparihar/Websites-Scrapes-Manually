import openpyxl
import urllib.parse
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\💻 Coding Websites.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
wb_f = openpyxl.load_workbook(fpath, data_only=False)

def clean_url(u):
    if not u: return ''
    u = str(u).strip().strip('\'"<>[]()')
    if ' ' in u or '\n' in u or '\t' in u: return ''
    if not u.startswith('http://') and not u.startswith('https://'):
        if '.' in u and not u.startswith('='): u = 'https://' + u
    parsed = urllib.parse.urlparse(u)
    if not parsed.netloc: return ''
    return u.rstrip('/')

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    links = []
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            cell = ws.cell(r, c)
            cell_f = ws_f.cell(r, c)
            hl = cell.hyperlink.target if cell.hyperlink else None
            val = str(cell.value or '').strip()
            u = ''
            if hl: u = hl
            elif val.startswith('http'): u = val
            elif cell_f.value and str(cell_f.value).startswith('=HYPERLINK'):
                m = re.search(r'=HYPERLINK\("?([^",\)]+)"?', str(cell_f.value))
                if m: u = m.group(1).strip()
            if u:
                cu = clean_url(u)
                if cu and 'chatgpt.com' not in cu:
                    links.append((r, c, val, cu))
    print(f'{sname:35s}: {len(links)} real non-chatgpt links')
