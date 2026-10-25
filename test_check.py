import os
import sys
import glob
import re
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

downloads_dir = r"C:\Users\satyanarayan verma\Downloads"
files = glob.glob(os.path.join(downloads_dir, "*.xlsx"))

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip().strip("'\"<>[]()")
    if " " in u or "\n" in u or "\t" in u:
        return ""
    if not u.startswith("http://") and not u.startswith("https://"):
        if "." in u and not u.startswith("="):
            u = "https://" + u
    return u.rstrip('/')

items = []

for f in files:
    fname = os.path.basename(f)
    wb = openpyxl.load_workbook(f, data_only=True)
    wb_f = openpyxl.load_workbook(f, data_only=False)
    for sname in wb.sheetnames:
        ws = wb[sname]
        ws_f = wb_f[sname]
        print(f"Checking {fname} -> Sheet: {sname} (Rows: {ws.max_row})")

print("Done checking.")
