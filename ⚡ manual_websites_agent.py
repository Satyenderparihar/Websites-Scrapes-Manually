"""
🌐 Autonomous Manual Websites Agent & Curated Directory Suite
100% Free, Zero-API-Fee, Autonomous Link Ingestion, Categorization, Deduplication & GitHub Auto-Sync Engine.
Designed for Antigravity & User: Satyender Parihar (Satyenderparihar)
"""

import os
import sys
import re
import json
import glob
import time
import subprocess
import urllib.parse
import requests
import urllib3
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "💾 websites_database.json")
MASTER_MD = os.path.join(BASE_DIR, "📝 Master_Curated_Websites_Report.md")
MASTER_XLSX = os.path.join(BASE_DIR, "📊 Websites_Master_List.xlsx")
README_FILE = os.path.join(BASE_DIR, "README.md")
INPUT_FILE = os.path.join(BASE_DIR, "💾 paste_links_here.txt")
DOWNLOADS_DIR = r"C:\Users\satyanarayan verma\Downloads"

# Standard Categories with Emojis and Descriptions
CATEGORIES = {
    "🎵 Music & Audio Streaming": {
        "folder": "📁 🎵 Music & Audio Streaming",
        "desc": "Official music streaming platforms, Indian & global audio hubs, vocal extraction and production tools."
    },
    "📻 Online Radio & Interactive Maps": {
        "folder": "📁 📻 Online Radio & Interactive Maps",
        "desc": "Global live interactive radio maps, globe rotating tuners, and international frequency streams."
    },
    "📺 Global TV & Media Discovery": {
        "folder": "📁 📺 Global TV & Media Discovery",
        "desc": "Worldwide live TV channels, vintage nostalgic television simulators, webcams & broadcast discovery."
    },
    "⬇️ Video & Media Downloaders": {
        "folder": "📁 ⬇️ Video & Media Downloaders",
        "desc": "Fast online downloaders & converters for YouTube, Instagram, Facebook, Pinterest, Twitter & video media."
    },
    "🤖 AI & Machine Learning Tools": {
        "folder": "📁 🤖 AI & Machine Learning Tools",
        "desc": "Cutting-edge Generative AI builders, neural models, agentic workflows, LLM tools & prompt suites."
    },
    "🎓 Free Courses & Learning Platforms": {
        "folder": "📁 🎓 Free Courses & Learning Platforms",
        "desc": "High-impact free certification courses, computer science tutorials, AI academies & skill bootcamps."
    },
    "📄 PDF & Document Utilities": {
        "folder": "📁 📄 PDF & Document Utilities",
        "desc": "Online PDF editors, format converters, document mergers, splitters and lightweight office utilities."
    },
    "🛠️ Web Utilities & Productivity": {
        "folder": "📁 🛠️ Web Utilities & Productivity",
        "desc": "Universal web tools, personal workspaces, open file search engines, and daily browser utilities."
    },
    "🎨 Design & Creator Tools": {
        "folder": "📁 🎨 Design & Creator Tools",
        "desc": "Social media template packs, 3D modeling, lightweight JS animation libraries, and graphic assets."
    },
    "🎮 Gaming & Entertainment": {
        "folder": "📁 🎮 Gaming & Entertainment",
        "desc": "Browser multiplayer party games, retro simulators, and casual entertainment platforms."
    },
    "🌐 General Curated Tools": {
        "folder": "📁 🌐 General Curated Tools",
        "desc": "Curated online resources and verified web applications across diverse domains."
    }
}

VALID_TLDS = set([
    'com', 'org', 'net', 'edu', 'gov', 'mil', 'io', 'ai', 'co', 'app', 'dev', 'tv', 'me', 'in', 'is', 'so', 'gg',
    'cc', 'ly', 'to', 'la', 'fm', 'xyz', 'tech', 'online', 'site', 'store', 'club', 'pro', 'studio', 'tools', 'garden',
    'lol', 'live', 'page', 'space', 'fun', 'design', 'world', 'media', 'stream', 'cloud'
])

def clean_url(u):
    if not u:
        return ""
    u = str(u).strip().strip("'\"<>[]()")
    if " " in u or "\n" in u or "\t" in u:
        return ""
    if not u.startswith("http://") and not u.startswith("https://"):
        if "." in u and not u.startswith("="):
            u = "https://" + u
    parsed = urllib.parse.urlparse(u)
    if not parsed.netloc:
        return ""
    
    parts = parsed.netloc.lower().split('.')
    if len(parts) < 2:
        return ""
    tld = parts[-1].split(':')[0]
    if tld not in VALID_TLDS and len(tld) > 6:
        return ""
        
    # Strip tracking parameters
    clean_query = []
    if parsed.query:
        for qp in parsed.query.split('&'):
            if not any(qp.startswith(t) for t in ['utm_', 'fbclid', 'gclid', 'ref_']):
                clean_query.append(qp)
    query_str = '&'.join(clean_query)
    clean = urllib.parse.urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, query_str, ''))
    return clean.rstrip('/')

def get_domain_key(url):
    parsed = urllib.parse.urlparse(url)
    domain = parsed.netloc.lower()
    if domain.startswith("www."):
        domain = domain[4:]
    path = parsed.path.rstrip('/')
    if len(path) > 1 and any(sub in domain for sub in ['youtube.com', 'instagram.com', 'github.com', 'coursera.org', 'analyticsvidhya.com', 'stpcomputereducation.com']):
        return f"{domain}{path}"
    return domain

def normalize_category(raw_cat, site_name="", url="", desc=""):
    text = f"{raw_cat} {site_name} {url} {desc}".lower()
    
    # 1. Radio
    if any(k in text for k in ['radio.garden', 'radio garden', 'online radio', 'radio map', 'live radio']):
        return "📻 Online Radio & Interactive Maps"
    
    # 2. TV
    if any(k in text for k in ['tv garden', 'tv.garden', 'famelack', 'live tv', 'tv discovery', 'myretrotvs', 'channel surfing']):
        return "📺 Global TV & Media Discovery"
        
    # 3. Music
    if any(k in text for k in ['music streaming', 'youtube music', 'spotify', 'soundcloud', 'gaana', 'jiosaavn', 'apple music', 'vocals', 'acapella', 'music production', 'online music']):
        return "🎵 Music & Audio Streaming"
        
    # 4. Downloaders
    if any(k in text for k in ['download', 'downloader', 'fdown', 'igram', 'ssyoutube', 'y2mate', 'cobalt', 'save', 'botdownloader', 'autosubtitles', 'video download']):
        return "⬇️ Video & Media Downloaders"
        
    # 5. Courses & Learning
    if any(k in text for k in ['course', 'academy', 'learn', 'education', 'tutorial', 'coursera', 'skillshop', 'training', 'beginner', 'study notes', 'learning']):
        return "🎓 Free Courses & Learning Platforms"
        
    # 6. PDF
    if any(k in text for k in ['pdf', 'sejda', 'tinywow']):
        return "📄 PDF & Document Utilities"
        
    # 7. AI
    if any(k in text for k in ['ai tools', 'replit', 'genai', 'agentic', 'deepseek', 'gpt', 'model', 'machine learning', 'artificial intelligence', 'ai builder']):
        return "🤖 AI & Machine Learning Tools"
        
    # 8. Design & Creator
    if any(k in text for k in ['design', 'creator tools', 'template', '3dtuning', 'animejs', 'animation', 'ui/ux', 'customization']):
        return "🎨 Design & Creator Tools"
        
    # 9. Gaming
    if any(k in text for k in ['gaming', 'game', 'airconsole', 'party game', 'multiplayer']):
        return "🎮 Gaming & Entertainment"
        
    # 10. Utility & Productivity
    if any(k in text for k in ['utility', 'workspace', 'productivity', 'tool', 'xtiles', '123apps', 'filepursuit', 'image.pi7.org']):
        return "🛠️ Web Utilities & Productivity"
        
    return "🌐 General Curated Tools"

def normalize_pricing(p):
    if not p:
        return "Free"
    p_lower = str(p).lower().strip()
    if "free" in p_lower and "paid" in p_lower:
        return "Free + Paid"
    if "freemium" in p_lower:
        return "Freemium"
    if "paid" in p_lower:
        return "Paid"
    if "free" in p_lower or "open source" in p_lower or "audit" in p_lower:
        return "Free"
    return str(p).strip()

def fetch_metadata(url):
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }
    try:
        resp = requests.get(url, headers=headers, timeout=3, verify=False)
        if resp.status_code == 200:
            soup = BeautifulSoup(resp.content, 'html.parser')
            title = soup.title.string.strip() if (soup.title and soup.title.string) else ""
            desc = ""
            meta_desc = soup.find('meta', attrs={'name': 'description'}) or soup.find('meta', attrs={'property': 'og:description'})
            if meta_desc and meta_desc.get('content'):
                desc = meta_desc['content'].strip()
            if title:
                title = re.sub(r'[\r\n\t]+', ' ', title).strip()
            if desc:
                desc = re.sub(r'[\r\n\t]+', ' ', desc).strip()
            return title, desc
    except Exception:
        pass
    return "", ""

def load_database():
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_database(data):
    with open(DB_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

def add_or_update_website(db, url, name="", category="", pricing="", description="", source="Manual", auto_fetch=False):
    cleaned = clean_url(url)
    if not cleaned:
        return None
        
    key = get_domain_key(cleaned)
    final_cat = normalize_category(category, name, cleaned, description)
    final_pricing = normalize_pricing(pricing)
    
    if key in db:
        rec = db[key]
        if name and (len(name) > len(rec.get("name", "")) or rec.get("name") == rec.get("key")):
            rec["name"] = name
        if final_cat != "🌐 General Curated Tools":
            rec["category"] = final_cat
        if final_pricing != "Free" or rec.get("pricing") == "Free":
            rec["pricing"] = final_pricing
        if description and len(description) > len(rec.get("description", "")):
            rec["description"] = description
        rec["url"] = cleaned
        rec["updated_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
        return rec
    else:
        meta_title, meta_desc = "", ""
        if auto_fetch and (not name or name == key):
            meta_title, meta_desc = fetch_metadata(cleaned)
            
        final_name = name if (name and name != key) else (meta_title[:50] if meta_title else key)
        final_desc = description if description else (meta_desc if meta_desc else f"Official platform resource for {final_name}")
        
        if final_cat == "🌐 General Curated Tools":
            final_cat = normalize_category(category, final_name, cleaned, final_desc)
            
        rec = {
            "key": key,
            "name": final_name,
            "url": cleaned,
            "category": final_cat,
            "pricing": final_pricing,
            "description": final_desc,
            "source": source,
            "added_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S")
        }
        db[key] = rec
        return rec

def parse_user_text_input(text):
    lines = text.strip().splitlines()
    parsed_items = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if any(h in line.lower() for h in ['website type', 'website name', 'official url', 'free / paid', 'link', 'category']):
            if 'http' not in line:
                continue

        if '\t' in line:
            parts = [p.strip() for p in line.split('\t') if p.strip()]
        elif '|' in line:
            parts = [p.strip() for p in line.split('|') if p.strip()]
        else:
            parts = [line]

        url, name, cat, pricing = "", "", "", ""
        for p in parts:
            if 'http://' in p or 'https://' in p or ('.' in p and '/' in p):
                m = re.search(r'(https?://[^\s]+)', p)
                if m:
                    url = m.group(1)
                else:
                    url = p
            elif any(pr in p.lower() for pr in ['free', 'paid', 'freemium', '$']):
                pricing = p
            elif any(c in p.lower() for c in ['streaming', 'music', 'radio', 'tv', 'download', 'ai', 'course', 'tools', 'utility', 'pdf', 'design']):
                cat = p
            else:
                if not name:
                    name = p
                elif not cat:
                    cat = p

        if len(parts) >= 4 and not url:
            for p in parts:
                if 'http' in p:
                    url = p
        if len(parts) == 4 and ('http' in parts[2]):
            cat, name, url, pricing = parts[0], parts[1], parts[2], parts[3]

        if url:
            parsed_items.append({
                "url": url,
                "name": name,
                "category": cat,
                "pricing": pricing,
                "description": ""
            })
    return parsed_items

def ingest_downloads_excel_files(db):
    files = glob.glob(os.path.join(DOWNLOADS_DIR, "*.xlsx"))
    total_added = 0
    print(f"\n📂 Scanning Downloads folder: Found {len(files)} Excel files...")
    
    for f in files:
        fname = os.path.basename(f)
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            wb_f = openpyxl.load_workbook(f, data_only=False)
            file_count = 0
            for sname in wb.sheetnames:
                ws = wb[sname]
                ws_f = wb_f[sname]
                
                header_row = 1
                headers = []
                for r in range(1, min(ws.max_row + 1, 15)):
                    row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
                    joined = " ".join(row_vals).lower()
                    if any(k in joined for k in ['website', 'url', 'link', 'course name', 'resource name']):
                        header_row = r
                        headers = [str(ws.cell(r, c).value or '').strip().lower() for c in range(1, ws.max_column + 1)]
                        break
                
                for r in range(header_row + 1, ws.max_row + 1):
                    name, url, cat, pricing, desc = "", "", "", "", ""
                    for c in range(1, ws.max_column + 1):
                        cell = ws.cell(r, c)
                        cell_f = ws_f.cell(r, c)
                        if cell.hyperlink and cell.hyperlink.target:
                            url = cell.hyperlink.target
                        elif cell_f.value and str(cell_f.value).startswith("=HYPERLINK"):
                            m = re.search(r'=HYPERLINK\("?([^",\)]+)"?', str(cell_f.value))
                            if m:
                                ref = m.group(1).strip()
                                if ref.startswith("http"):
                                    url = ref
                                else:
                                    try:
                                        url = ws[ref].value
                                    except Exception:
                                        pass
                    
                    for c, h in enumerate(headers, 1):
                        val = str(ws.cell(r, c).value or "").strip()
                        if not val or val == "None":
                            continue
                        if any(k in h for k in ['url', 'link']) or ('website' in h and ('.' in val or 'http' in val)):
                            if not url or ('http' not in url and 'http' in val):
                                url = val
                        elif any(k in h for k in ['name', 'resource', 'course']):
                            if not name:
                                name = val
                        elif any(k in h for k in ['type', 'category']):
                            if not cat:
                                cat = val
                        elif any(k in h for k in ['free', 'pricing', 'paid']):
                            if not pricing:
                                pricing = val
                        elif 'desc' in h:
                            if not desc:
                                desc = val
                    
                    if not url:
                        for c in range(1, ws.max_column + 1):
                            v = str(ws.cell(r, c).value or "").strip()
                            if v.startswith("http://") or v.startswith("https://"):
                                url = v
                                break
                    
                    if not name and url:
                        for c in range(1, ws.max_column + 1):
                            v = str(ws.cell(r, c).value or "").strip()
                            if v and v != url and v != "Click Here" and len(v) > 2:
                                name = v
                                break

                    cleaned = clean_url(url)
                    if cleaned:
                        rec = add_or_update_website(db, cleaned, name=name, category=cat or sname, pricing=pricing, description=desc, source=f"Excel: {fname}", auto_fetch=False)
                        if rec:
                            file_count += 1
                            total_added += 1
            print(f"  ✓ {fname} -> Ingested {file_count} records.")
        except Exception as e:
            print(f"  ⚠️ Error processing {fname}: {e}")
            
    print(f"✅ Total processed from Downloads: {total_added} records.")
    return total_added

def build_category_directories_and_reports(db):
    by_category = {cat: [] for cat in CATEGORIES.keys()}
    for key, rec in db.items():
        cat = rec.get("category", "🌐 General Curated Tools")
        if cat not in by_category:
            by_category[cat] = []
        by_category[cat].append(rec)
        
    for cat_name, items in by_category.items():
        items.sort(key=lambda x: x.get("name", "").lower())
        
    # 1. Create Dedicated Category Subfolders and Reports
    for cat_name, meta in CATEGORIES.items():
        folder_path = os.path.join(BASE_DIR, meta["folder"])
        os.makedirs(folder_path, exist_ok=True)
        cat_items = by_category.get(cat_name, [])
        cat_safe = re.sub(r'[^\w\s-]', '', cat_name).strip().replace(' ', '_')
        report_file = os.path.join(folder_path, f"📝 {cat_safe}_Report.md")
        
        md_lines = [
            f"# {cat_name}",
            "",
            f"> {meta['desc']}  ",
            f"> **Total Curated Tools**: `{len(cat_items)}` | **Zero Repeats & Deduplicated**",
            "",
            "---",
            "",
            "| # | Tool / Website Name | Official Link | Pricing | Description |",
            "| :-: | :--- | :--- | :-: | :--- |"
        ]
        
        for i, item in enumerate(cat_items, 1):
            name = (item.get("name") or item.get("key")).replace('|', '-')
            url = item.get("url")
            pricing = item.get("pricing", "Free")
            desc = (item.get("description") or f"Official platform for {name}").replace('|', '-')
            md_lines.append(f"| **{i}** | **{name}** | [Visit Website]({url}) | `{pricing}` | {desc} |")
            
        md_lines.append("\n---\n*Auto-generated & Curated by Autonomous Websites Agent.*")
        with open(report_file, 'w', encoding='utf-8') as f:
            f.write("\n".join(md_lines))

    # 2. Build Master Markdown Report
    master_lines = [
        "# 🌐 Curated Websites & Tools Master Directory",
        "",
        "> **Consolidated, Deduplicated & Isolated Directory of High-Utility Web Platforms**  ",
        f"> **Total Unique Websites**: `{len(db)}` across `{len(CATEGORIES)}` distinct categories.  ",
        "> **Repository**: [https://github.com/Satyenderparihar/Websites-Scrapes-Manually](https://github.com/Satyenderparihar/Websites-Scrapes-Manually)",
        "",
        "---",
        "",
        "## 📑 Quick Category Navigation",
        ""
    ]
    
    for cat_name, meta in CATEGORIES.items():
        count = len(by_category.get(cat_name, []))
        folder_link = urllib.parse.quote(meta["folder"])
        master_lines.append(f"- [{cat_name}]({folder_link}) — `{count} Websites` — *{meta['desc']}*")
        
    master_lines.append("\n---\n")
    
    for cat_name, meta in CATEGORIES.items():
        items = by_category.get(cat_name, [])
        if not items:
            continue
        master_lines.extend([
            f"## {cat_name}",
            "",
            f"*{meta['desc']}* (Total: `{len(items)}`)",
            "",
            "| # | Website Name | Official URL | Pricing | Description |",
            "| :-: | :--- | :--- | :-: | :--- |"
        ])
        for i, item in enumerate(items, 1):
            name = (item.get("name") or item.get("key")).replace('|', '-')
            url = item.get("url")
            pricing = item.get("pricing", "Free")
            desc = (item.get("description") or f"Official platform for {name}").replace('|', '-')
            master_lines.append(f"| **{i}** | **{name}** | [{url}]({url}) | `{pricing}` | {desc} |")
        master_lines.append("\n---\n")

    with open(MASTER_MD, 'w', encoding='utf-8') as f:
        f.write("\n".join(master_lines))

    # 3. Build Root README.md
    readme_lines = [
        "# 🌐 Curated Useful Websites & Media Tools Directory",
        "",
        "![Status](https://img.shields.io/badge/Status-Active%20%26%20Auto--Synced-success) ",
        f"![Total Websites](https://img.shields.io/badge/Total%20Websites-{len(db)}-blue) ",
        f"![Categories](https://img.shields.io/badge/Categories-{len(CATEGORIES)}-orange) ",
        "![License](https://img.shields.io/badge/License-MIT-green)",
        "",
        "> **100% Free, Curated & Deduplicated Directory of Useful Websites, Media Streaming & Web Utilities.**  ",
        "> Maintained separately from automated scrapers to ensure curated accuracy and zero clutter.",
        "",
        "---",
        "",
        "## 📚 Master Directory & Download Links",
        "",
        "- 📖 **[View Complete Master Categorized Report (Markdown)](📝%20Master_Curated_Websites_Report.md)**",
        "- 📊 **[Download Formatted Master Excel Workbook](📊%20Websites_Master_List.xlsx)**",
        "- 💾 **[Raw Structured Database (JSON)](💾%20websites_database.json)**",
        "",
        "---",
        "",
        "## 🚀 Dedicated Categories Breakdown",
        "",
        "| Category | Count | Dedicated Folder & Report | Description |",
        "| :--- | :---: | :--- | :--- |"
    ]
    
    for cat_name, meta in CATEGORIES.items():
        count = len(by_category.get(cat_name, []))
        folder_link = urllib.parse.quote(meta["folder"])
        cat_safe = re.sub(r'[^\w\s-]', '', cat_name).strip().replace(' ', '_')
        report_link = f"{folder_link}/📝%20{cat_safe}_Report.md"
        readme_lines.append(f"| **{cat_name}** | **{count}** | [📁 Explore Category]({folder_link}) / [📖 Report]({report_link}) | {meta['desc']} |")
        
    readme_lines.extend([
        "",
        "---",
        "",
        "## ⚡ 1-Click Autonomous Agent (How It Works)",
        "",
        "Whenever you have new website links, tables, or notes:",
        "1. Double click **`⚡ Run Manual Websites Agent.bat`** (or use the Desktop Shortcut **`⚡ Manual Websites Agent.lnk`**).",
        "2. Paste your links or text (or put them in `💾 paste_links_here.txt`).",
        "3. The agent automatically:",
        "   - Fetches web titles & meta descriptions,",
        "   - Classifies them into the appropriate category,",
        "   - Deduplicates and prevents any duplicates (`Zero Repeats`),",
        "   - Updates the Markdown reports and multi-sheet Excel file,",
        "   - **Automatically commits & pushes to GitHub in 1 click!**",
        "",
        "---",
        "**Maintained with ❤️ by Satyender Parihar & Antigravity Autonomous Agent.**"
    ])
    
    with open(README_FILE, 'w', encoding='utf-8') as f:
        f.write("\n".join(readme_lines))

    # 4. Build Multi-Sheet Formatted Excel File
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    center_align = Alignment(horizontal='center', vertical='center')

    # Master Sheet
    ws_master = wb.create_sheet(title="All Websites Master")
    headers = ["#", "Website Name", "Category", "Official URL", "Pricing", "Description", "Source"]
    ws_master.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_master.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    row_num = 2
    for cat_name, items in by_category.items():
        for item in items:
            ws_master.append([
                row_num - 1,
                item.get("name") or item.get("key"),
                item.get("category"),
                item.get("url"),
                item.get("pricing", "Free"),
                item.get("description", ""),
                item.get("source", "Manual")
            ])
            ws_master.cell(row_num, 1).alignment = center_align
            ws_master.cell(row_num, 1).font = data_font
            ws_master.cell(row_num, 2).font = Font(name="Segoe UI", size=10, bold=True)
            ws_master.cell(row_num, 3).font = data_font
            url_cell = ws_master.cell(row_num, 4)
            url_cell.font = link_font
            url_cell.hyperlink = item.get("url")
            ws_master.cell(row_num, 5).alignment = center_align
            ws_master.cell(row_num, 5).font = data_font
            ws_master.cell(row_num, 6).font = data_font
            ws_master.cell(row_num, 7).font = data_font
            for c in range(1, 8):
                ws_master.cell(row_num, c).border = border_thin
            row_num += 1

    for col in ws_master.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws_master.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

    # Individual Category Sheets
    for cat_name, items in by_category.items():
        if not items:
            continue
        short_title = re.sub(r'[^\w\s]', '', cat_name)[:30].strip()
        ws_cat = wb.create_sheet(title=short_title)
        cat_headers = ["#", "Website Name", "Official URL", "Pricing", "Description"]
        ws_cat.append(cat_headers)
        for col_idx in range(1, len(cat_headers) + 1):
            cell = ws_cat.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            
        for i, item in enumerate(items, 1):
            r = i + 1
            ws_cat.append([
                i,
                item.get("name") or item.get("key"),
                item.get("url"),
                item.get("pricing", "Free"),
                item.get("description", "")
            ])
            ws_cat.cell(r, 1).alignment = center_align
            ws_cat.cell(r, 1).font = data_font
            ws_cat.cell(r, 2).font = Font(name="Segoe UI", size=10, bold=True)
            u_cell = ws_cat.cell(r, 3)
            u_cell.font = link_font
            u_cell.hyperlink = item.get("url")
            ws_cat.cell(r, 4).alignment = center_align
            ws_cat.cell(r, 4).font = data_font
            ws_cat.cell(r, 5).font = data_font
            for c in range(1, 6):
                ws_cat.cell(r, c).border = border_thin

        for col in ws_cat.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws_cat.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

    wb.save(MASTER_XLSX)
    print(f"📊 Excel file saved: {MASTER_XLSX}")
    print(f"📝 Master Markdown Report saved: {MASTER_MD}")
    print(f"📄 README.md updated: {README_FILE}")

def git_commit_and_push(commit_msg):
    try:
        subprocess.run(["git", "add", "."], cwd=BASE_DIR, check=True)
        status = subprocess.run(["git", "status", "--porcelain"], cwd=BASE_DIR, capture_output=True, text=True)
        if not status.stdout.strip():
            print("ℹ️ No new changes to commit in git.")
            return True
        subprocess.run(["git", "commit", "-m", commit_msg], cwd=BASE_DIR, check=True)
        print("🚀 Pushing to GitHub...")
        subprocess.run(["git", "push", "-u", "origin", "main"], cwd=BASE_DIR, check=True)
        print("✅ Successfully pushed to GitHub (https://github.com/Satyenderparihar/Websites-Scrapes-Manually)!")
        return True
    except Exception as e:
        print(f"⚠️ Git operation error: {e}")
        return False

def main():
    db = load_database()
    print("=" * 60)
    print("🌐 Autonomous Manual Websites Agent & Curated Hub")
    print(f"📁 Working Directory: {BASE_DIR}")
    print(f"💾 Current Database Count: {len(db)} websites")
    print("=" * 60)
    
    if len(sys.argv) > 1:
        arg = sys.argv[1].lower()
        if arg in ["--all", "--sync", "--rebuild"]:
            build_category_directories_and_reports(db)
            git_commit_and_push(f"feat: rebuild and sync curated directory ({len(db)} websites)")
            return
        elif arg == "--ingest-downloads":
            ingest_downloads_excel_files(db)
            save_database(db)
            build_category_directories_and_reports(db)
            git_commit_and_push(f"feat: ingested websites from Downloads excel files ({len(db)} websites)")
            return

    while True:
        print("\nSelect an action:")
        print("  [1] ➕ Add / Paste Website(s) or Table (Multi-line supported)")
        print("  [2] 📄 Process Input File (💾 paste_links_here.txt)")
        print("  [3] 📥 Import All Excel Files from Downloads")
        print("  [4] 🔄 Rebuild Reports, Excel & Push to GitHub")
        print("  [5] 🌐 Open GitHub Repository in Browser")
        print("  [6] 📊 Open Master Excel File")
        print("  [0] 🚪 Exit")
        
        choice = input("\nEnter choice [0-6]: ").strip()
        
        if choice == "1":
            print("\nPaste your websites below (table, URL list, or details).")
            print("When finished, press Enter on an empty line:")
            lines = []
            while True:
                try:
                    l = input()
                    if not l.strip():
                        break
                    lines.append(l)
                except EOFError:
                    break
            raw_text = "\n".join(lines)
            if raw_text.strip():
                items = parse_user_text_input(raw_text)
                print(f"Found {len(items)} websites to process...")
                for item in items:
                    rec = add_or_update_website(db, item["url"], name=item["name"], category=item["category"], pricing=item["pricing"], description=item["description"], source="Direct Input", auto_fetch=True)
                    if rec:
                        print(f"  ✓ Added: {rec['name']} -> {rec['category']} ({rec['pricing']})")
                save_database(db)
                build_category_directories_and_reports(db)
                git_commit_and_push(f"feat: added {len(items)} curated websites (Total: {len(db)})")
            else:
                print("No input provided.")
                
        elif choice == "2":
            if not os.path.exists(INPUT_FILE):
                with open(INPUT_FILE, 'w', encoding='utf-8') as f:
                    f.write("# Paste website URLs, markdown tables, or text here.\n# One per line or tab-separated.\n")
                print(f"Created template file: {INPUT_FILE}. Paste links there and run again.")
            else:
                with open(INPUT_FILE, 'r', encoding='utf-8') as f:
                    raw_text = f.read()
                items = parse_user_text_input(raw_text)
                if items:
                    print(f"Found {len(items)} websites in {INPUT_FILE}...")
                    for item in items:
                        rec = add_or_update_website(db, item["url"], name=item["name"], category=item["category"], pricing=item["pricing"], description=item["description"], source="Input File", auto_fetch=True)
                        if rec:
                            print(f"  ✓ Added: {rec['name']} -> {rec['category']}")
                    save_database(db)
                    build_category_directories_and_reports(db)
                    git_commit_and_push(f"feat: synced {len(items)} websites from input file (Total: {len(db)})")
                else:
                    print("No valid websites found in file.")
                    
        elif choice == "3":
            ingest_downloads_excel_files(db)
            save_database(db)
            build_category_directories_and_reports(db)
            git_commit_and_push(f"feat: imported websites from Downloads Excel files (Total: {len(db)})")
            
        elif choice == "4":
            build_category_directories_and_reports(db)
            git_commit_and_push(f"feat: rebuild and sync master directory (Total: {len(db)} websites)")
            
        elif choice == "5":
            os.system("start https://github.com/Satyenderparihar/Websites-Scrapes-Manually")
            
        elif choice == "6":
            if os.path.exists(MASTER_XLSX):
                os.startfile(MASTER_XLSX)
            else:
                print("Excel file not created yet.")
                
        elif choice == "0":
            print("Exiting. Bye!")
            break

if __name__ == "__main__":
    main()
