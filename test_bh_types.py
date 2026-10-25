import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_power = r"C:\Users\satyanarayan verma\Downloads\🌐 Powerful Websites.xlsx"
wb = openpyxl.load_workbook(f_power, data_only=True)
ws = wb['Websites - BH']

def derive_website_type(site, cat, desc):
    d = f"{site} {cat} {desc}".lower()
    if cat == "AI Tools":
        if 'builder' in d or 'replit' in d: return "AI Website Builder & IDE"
        if 'vocal' in d or 'acapella' in d: return "AI Vocal Extractor"
        if 'video' in d: return "AI Video Generator & Pre-Production"
        if 'audio' in d or 'voice' in d: return "AI Audio & Voice Tool"
        if 'image' in d or 'photo' in d: return "AI Image Generator & Enhancer"
        if 'writing' in d or 'text' in d: return "AI Writing & Copywriting Assistant"
        return "AI Productivity & Generation Tool"
    elif cat == "Gaming":
        if 'multiplayer' in d or 'phone' in d: return "Phone-Controller Multiplayer Games"
        return "Browser Game & Entertainment"
    elif cat in ["Creator Tools", "Creative", "Design"]:
        if 'template' in d: return "Social Media Template Booster"
        if 'car' in d or '3d' in d or 'tuning' in d: return "3D Virtual Vehicle Customizer"
        if 'animation' in d: return "JavaScript Web Animation Engine"
        if 'color' in d or 'palette' in d: return "Color Palette & Visual Design Tool"
        if 'font' in d or 'typography' in d: return "Typography & Font Tool"
        if 'icon' in d or 'logo' in d: return "Icon & Logo Asset Library"
        return f"{cat} Platform"
    elif cat in ["Utility", "Productivity", "Developer"]:
        if 'car repair' in d or 'vehicle' in d: return "DIY Vehicle Repair Guides"
        if 'tv' in d or 'retro' in d: return "Retro TV Channel Simulator"
        if 'archive' in d or 'permanent' in d: return "Web Archiving Utility"
        if 'repair' in d or 'fixit' in d: return "Hardware & Device Repair Guides"
        if 'workspace' in d or 'notes' in d: return "Workspace & Productivity Tool"
        if 'code' in d or 'developer' in d or 'api' in d: return "Developer & Code Utility"
        if 'file' in d or 'convert' in d: return "File Converter & Search Engine"
        return f"{cat} Tool"
    elif cat == "Education":
        return "Educational Learning Platform"
    return "Curated Web Tool"

for r in range(2, 20):
    site = str(ws.cell(r, 1).value or '')
    desc = str(ws.cell(r, 2).value or '')
    cat = str(ws.cell(r, 3).value or '')
    wtype = derive_website_type(site, cat, desc)
    print(f"[{cat}] {site} -> Type: {wtype}")
