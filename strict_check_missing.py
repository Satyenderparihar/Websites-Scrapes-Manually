import openpyxl
import json
import sys
import urllib.parse
import re

sys.stdout.reconfigure(encoding='utf-8')

fpath = r'C:\Users\satyanarayan verma\Downloads\💻 Coding Websites.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
wb_f = openpyxl.load_workbook(fpath, data_only=False)

with open(r'D:\Antigravity - Agent\🌐 Websites Scrapes Manually\💾 websites_database.json', encoding='utf-8') as f:
    db = json.load(f)

# Normalize all names and URLs in DB
db_names = set(v['name'].lower().strip() for v in db.values())
db_urls = set(v['url'].lower().rstrip('/') for v in db.values())
db_keys = set(db.keys())

# Also collect words from DB names
db_words = set()
for n in db_names:
    db_words.add(n)
    for part in re.split(r'[\s\-_–—]+', n):
        if len(part) > 2:
            db_words.add(part)

missing_per_sheet = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    missing = []
    
    if sname in ["All Project Ideas", "C Project Ideas"]:
        for r in range(1, ws.max_row + 1):
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2 = str(ws.cell(r, 2).value or '').strip()
            pname = ""
            if sname == "All Project Ideas":
                if c2 and c2 not in ["Project Title", "Link", "Click Here"] and len(c2) > 3:
                    pname = c2
                elif c1 and not any(k in c1 for k in ["Category", "Projects", "Project Ideas"]) and len(c1) > 3:
                    pname = c1
            else:
                if c1 and not any(k in c1 for k in ["Project Name", "Project Ideas", "C Project", "C++ Project", "Java Project", "Python Project"]) and len(c1) > 3:
                    pname = c1
            if pname and len(pname) > 3:
                # check if in db
                p_norm = pname.lower().strip()
                found = any(p_norm in n or n in p_norm for n in db_names)
                if not found:
                    missing.append((r, pname))
                    
    elif sname == "Y-I-T - BH":
        for r in range(2, ws.max_row + 1):
            c1 = str(ws.cell(r, 1).value or '').strip()
            c2 = str(ws.cell(r, 2).value or '').strip()
            if c2 and c2 not in ["Channel Name", "Link", "None"] and len(c2) > 2:
                # creator check
                c_norm = c2.lower().strip()
                found = any(c_norm in n or any(w in n for w in c_norm.split() if len(w) > 3) for n in db_names)
                if not found:
                    missing.append((r, c2))
                    
    elif sname == "Pythons - BH":
        for r in range(2, ws.max_row + 1):
            c2 = str(ws.cell(r, 2).value or '').strip()
            if c2 and c2 not in ["Website", "None"] and len(c2) > 2:
                c_norm = c2.lower().strip()
                found = any(c_norm in n or any(w in n for w in c_norm.split() if len(w) > 3) for n in db_names)
                if not found:
                    missing.append((r, c2))
                    
    elif sname == "Sheet32 - BH":
        for r in range(2, ws.max_row + 1):
            for c in range(1, 9):
                cell = ws.cell(r, c)
                val = str(cell.value or '').strip()
                hl = cell.hyperlink.target if cell.hyperlink else None
                if val and val not in ["Coding Platforms", "None", "Link", "Click Here"] and len(val) > 1:
                    v_norm = val.lower().strip()
                    hl_norm = hl.lower().rstrip('/') if hl else ""
                    found = (v_norm in db_names) or (hl_norm in db_urls) or any(v_norm in n for n in db_names)
                    if not found and hl:
                        # check domain in keys
                        domain = urllib.parse.urlparse(hl).netloc.replace('www.', '').lower()
                        found = any(domain in k for k in db_keys)
                    if not found and not val.startswith("="):
                        missing.append((r, c, val, hl))
                        
    else: # standard sheets
        for r in range(2, ws.max_row + 1):
            name = str(ws.cell(r, 2).value or '').strip()
            cell_3 = ws.cell(r, 3)
            hl = cell_3.hyperlink.target if cell_3.hyperlink else str(cell_3.value or '').strip()
            if name and name not in ["Website Name", "None"] and len(name) > 1:
                n_norm = name.lower().strip()
                hl_norm = hl.lower().rstrip('/') if hl else ""
                found = (n_norm in db_names) or (hl_norm in db_urls) or any(n_norm in n for n in db_names)
                if not found:
                    missing.append((r, name, hl))
                    
    missing_per_sheet[sname] = missing
    print(f"Sheet: {sname:35s} | Missing: {len(missing)}")

print("\n--- Summary of any missing entries ---")
for s, miss in missing_per_sheet.items():
    if miss:
        print(f"\nSheet {s}:")
        for m in miss[:10]:
            print("  ", m)
